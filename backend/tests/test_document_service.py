from pathlib import Path

from docx import Document
from reportlab.pdfgen import canvas

from app.services.document_service import extract_text


TEST_DIR = Path(__file__).parent


def create_test_docx():
    file_path = TEST_DIR / "sample_test.docx"

    document = Document()
    document.add_paragraph("Sovereign AI Workbench")
    document.add_paragraph("This is a test document for DOCX extraction.")
    document.save(file_path)

    return file_path


def create_test_pdf():
    file_path = TEST_DIR / "sample_test.pdf"

    pdf = canvas.Canvas(str(file_path))
    pdf.drawString(100, 750, "Sovereign AI Workbench")
    pdf.drawString(100, 730, "This is a test document for PDF extraction.")
    pdf.save()

    return file_path


def test_docx_extraction():
    file_path = create_test_docx()

    text = extract_text(str(file_path))

    assert "Sovereign AI Workbench" in text
    assert "DOCX extraction" in text


def test_pdf_extraction():
    file_path = create_test_pdf()

    text = extract_text(str(file_path))

    assert "Sovereign AI Workbench" in text
    assert "PDF extraction" in text
def test_chunk_text():
    from app.services.document_service import chunk_text

    text = "A" * 1200
    chunks = chunk_text(text, chunk_size=500, overlap=50)

    assert len(chunks) == 3
    assert len(chunks[0]) == 500
    assert len(chunks[1]) == 500
    assert len(chunks[2]) == 300


def test_chunk_text_empty():
    from app.services.document_service import chunk_text

    chunks = chunk_text("")

    assert chunks == []

def test_document_ingestion():
    from app.services.ingestion_service import ingest_document

    file_path = TEST_DIR / "sample_test.docx"

    chunk_count = ingest_document(str(file_path))

    assert chunk_count > 0